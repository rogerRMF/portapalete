# -*- coding: utf-8 -*-
import streamlit as st
import pandas as pd
import os
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment

CAMINHO_EXCEL = r"C:\Users\roger\OneDrive\Área de Trabalho\DESENHO PORTA PALETES.xlsx"
ABA_ESTOQUE = "Estoque CD"

# Mapeamento de endereços para células (incluindo os novos solicitados)
enderecos_celulas = {
    "75A": "D19",
    "73A": "F19",
    "71A": "I19",
    "69A": "L19",
    "67A": "O19",
    "65A": "R19",
    "75B": "C15",
    "73B": "F15",
    "71B": "I15",
    "69B": "L15",
    "67B": "O15",
    "65B": "R15",
    "75C": "C11",
    "73C": "F11",
    "71C": "I11",
    "69C": "L11",
    "67C": "O11",
    "65C": "R11",
    "75D": "C7",
    "73D": "F7",
    "71D": "I7",
    "69D": "L7",
    "67D": "O7",
    "65D": "R7",
    "75E": "C3",
    "73E": "F3",
    "71E": "I3",
    "69E": "L3",
    "67E": "O3",
    "65E": "R3"
}


def registrar_produto_na_celula(produto, endereco):
    if endereco not in enderecos_celulas:
        st.error(
            "Cadastro automático só implementado para os endereços cadastrados no sistema.")
        return

    celula = enderecos_celulas[endereco]
    wb = load_workbook(CAMINHO_EXCEL)
    ws = wb[ABA_ESTOQUE]

    # Formatação: tudo maiúsculo, quebra de linha, letra branca, fundo vermelho
    valor = produto.upper()
    fill = PatternFill(start_color="FF0000",
                       end_color="FF0000", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    alignment = Alignment(
        wrap_text=True, horizontal="center", vertical="center")

    # Se a célula for mesclada, escreva na célula superior esquerda do range mesclado
    for merged_range in ws.merged_cells.ranges:
        if celula in merged_range:
            cell = ws[merged_range.coord.split(":")[0]]
            break
    else:
        cell = ws[celula]

    cell.value = valor
    cell.fill = fill
    cell.font = font
    cell.alignment = alignment

    wb.save(CAMINHO_EXCEL)


def liberar_endereco(endereco):
    if endereco not in enderecos_celulas:
        st.error(
            "Saída automática só implementada para os endereços cadastrados no sistema.")
        return

    celula = enderecos_celulas[endereco]
    wb = load_workbook(CAMINHO_EXCEL)
    ws = wb[ABA_ESTOQUE]

    # Formatação: DISPONÍVEL, letra branca, fundo verde, quebra de linha
    valor = "DISPONÍVEL"
    fill = PatternFill(start_color="00B050",
                       end_color="00B050", fill_type="solid")
    font = Font(color="FFFFFF", bold=True)
    alignment = Alignment(
        wrap_text=True, horizontal="center", vertical="center")

    for merged_range in ws.merged_cells.ranges:
        if celula in merged_range:
            cell = ws[merged_range.coord.split(":")[0]]
            break
    else:
        cell = ws[celula]

    cell.value = valor
    cell.fill = fill
    cell.font = font
    cell.alignment = alignment

    wb.save(CAMINHO_EXCEL)


def endereco_disponivel(endereco):
    if endereco not in enderecos_celulas:
        return False  # Endereço não mapeado

    celula = enderecos_celulas[endereco]
    wb = load_workbook(CAMINHO_EXCEL, data_only=True)
    ws = wb[ABA_ESTOQUE]

    for merged_range in ws.merged_cells.ranges:
        if celula in merged_range:
            cell = ws[merged_range.coord.split(":")[0]]
            break
    else:
        cell = ws[celula]

    valor = str(cell.value).strip().upper() if cell.value else ""
    return valor == "" or valor == "DISPONÍVEL"


st.set_page_config(page_title="Gestão de Estoque",
                   page_icon="📦", layout="wide")

st.title("📦 Gestão de Estoque - Entradas e Saídas")

enderecos = list(enderecos_celulas.keys())

# Adicione uma chave para o estado do produto
if "produto_input" not in st.session_state:
    st.session_state.produto_input = ""

with st.form("form_estoque"):
    st.subheader("📋 Registrar Produto")

    acao = st.radio("Ação", ["Entrada", "Saída"], key="acao_radio")
    endereco = st.selectbox(
        "Endereço", options=enderecos, key="endereco_select")

    # Busca o valor atual na célula correspondente
    produto_atual = ""
    if endereco:
        celula = enderecos_celulas.get(endereco)
        if celula:
            wb = load_workbook(CAMINHO_EXCEL, data_only=True)
            ws = wb[ABA_ESTOQUE]
            for merged_range in ws.merged_cells.ranges:
                if celula in merged_range:
                    cell = ws[merged_range.coord.split(":")[0]]
                    break
            else:
                cell = ws[celula]
            produto_atual = str(cell.value) if cell.value else ""

    # Input para ambos, SEM bloqueio (disabled=False)
    if acao == "Entrada":
        produto = st.text_input(
            "Nome do Produto", value=st.session_state.produto_input, key="produto_text")
    else:  # Saída
        produto = st.text_input(
            "Produto no Endereço", value=produto_atual, key="produto_saida", disabled=False)

    submitted = st.form_submit_button("Registrar")

    if submitted:
        if acao == "Entrada":
            if produto.strip() == "":
                st.error("Digite o nome do produto.")
            elif not endereco_disponivel(endereco):
                st.error(
                    "Endereço OCUPADO! Use outro endereço disponível para cadastrar.")
                # Mostra os endereços disponíveis
                enderecos_disponiveis = [
                    e for e in enderecos if endereco_disponivel(e)]
                if enderecos_disponiveis:
                    st.info("Endereços disponíveis: " +
                            ", ".join(enderecos_disponiveis))
                else:
                    st.warning("Nenhum endereço disponível no momento.")
            else:
                registrar_produto_na_celula(produto, endereco)
                st.success(
                    f"Produto '{produto}' registrado no endereço {endereco}.")
                st.session_state.produto_input = ""
                st.rerun()
        elif acao == "Saída":
            liberar_endereco(endereco)
            st.success(
                f"Endereço {endereco} liberado e marcado como DISPONÍVEL.")
            st.session_state.produto_input = ""
            st.rerun()

st.info(
    "Cadastro e saída direto nas células D19 (75A), F19 (73A), I19 (71A), L19 (69A), O19 (67A), R19 (65A), "
    "C15 (75B), F15 (73B), I15 (71B), L15 (69B), O15 (67B), R15 (65B), "
    "C11 (75C), F11 (73C), I11 (71C), L11 (69C), O11 (67C), R11 (65C), "
    "C7 (75D), F7 (73D), I7 (71D), L7 (69D), O7 (67D), R7 (65D), "
    "C3 (75E), F3 (73E), I3 (71E), L3 (69E), O3 (67E), R3 (65E) da planilha."
)

st.subheader("📋 Endereços Ocupados")

# Carrega os dados atuais dos endereços
wb = load_workbook(CAMINHO_EXCEL, data_only=True)
ws = wb[ABA_ESTOQUE]
dados = []
for endereco, celula in enderecos_celulas.items():
    for merged_range in ws.merged_cells.ranges:
        if celula in merged_range:
            cell = ws[merged_range.coord.split(":")[0]]
            break
    else:
        cell = ws[celula]
    valor = str(cell.value) if cell.value else ""
    status = "Disponível" if valor.strip().upper(
    ) == "" or valor.strip().upper() == "DISPONÍVEL" else "Ocupado"
    dados.append({"Endereço": endereco, "Produto": valor, "Status": status})

df = pd.DataFrame(dados)


def colorir_status(val):
    if val == "Ocupado":
        return 'background-color: red; color: white'
    elif val == "Disponível":
        return 'background-color: green; color: white'
    return ''


st.dataframe(
    df.style.applymap(colorir_status, subset=["Status"]),
    use_container_width=True
)
